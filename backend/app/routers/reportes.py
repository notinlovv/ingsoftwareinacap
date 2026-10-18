import io

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import auth, database, models

router = APIRouter(prefix="/reportes", tags=["Reportes"])


@router.get("/avance.xlsx", dependencies=[Depends(auth.get_current_user)])
def exportar_reporte_avance_excel(
    linea_id: int | None = Query(
        default=None,
        description="ID de la línea de fondeo a exportar (si se omite, exporta todas las líneas)",
    ),
    db: Session = Depends(database.get_db),
):
    """
    Genera y devuelve un archivo Excel (.xlsx) usando openpyxl con el detalle
    de los registros de avance y una fila resumen con el total acumulado.
    """
    linea_filtrada = None
    if linea_id is not None:
        linea_filtrada = (
            db.query(models.LineaFondeo)
            .filter(models.LineaFondeo.id == linea_id)
            .first()
        )
        if not linea_filtrada:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Línea de fondeo con ID {linea_id} no encontrada",
            )

    query = db.query(models.RegistroAvance)
    if linea_id is not None:
        query = query.filter(models.RegistroAvance.linea_id == linea_id)

    registros = query.order_by(models.RegistroAvance.fecha.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Avance Líneas de Fondeo"

    # Estilos institucionales
    header_fill = PatternFill(
        start_color="1B365D", end_color="1B365D", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill(
        start_color="E8EEF5", end_color="E8EEF5", fill_type="solid"
    )
    total_font = Font(bold=True, color="1B365D", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )

    # Encabezados de la tabla
    encabezados = [
        "UUID Registro",
        "ID Línea",
        "Nombre Línea",
        "Supervisor (Email)",
        "Fecha Registro (UTC)",
        "Metros Avanzados (m)",
        "Latitud",
        "Longitud",
        "Estado",
        "Observación",
        "Sincronizado En (UTC)",
    ]
    ws.append(encabezados)

    for col_idx in range(1, len(encabezados) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    total_acumulado = 0.0

    for reg in registros:
        nombre_linea = reg.linea.nombre if reg.linea else f"Línea {reg.linea_id}"
        email_usuario = reg.usuario.email if reg.usuario else str(reg.usuario_id)
        fecha_str = reg.fecha.strftime("%Y-%m-%d %H:%M:%S") if reg.fecha else ""
        sync_str = (
            reg.sincronizado_en.strftime("%Y-%m-%d %H:%M:%S")
            if reg.sincronizado_en
            else ""
        )
        metros = float(reg.metros_avanzados or 0.0)
        total_acumulado += metros

        fila = [
            reg.id,
            reg.linea_id,
            nombre_linea,
            email_usuario,
            fecha_str,
            round(metros, 2),
            reg.latitud,
            reg.longitud,
            reg.estado,
            reg.observacion or "",
            sync_str,
        ]
        ws.append(fila)

        row_idx = ws.max_row
        for col_idx in range(1, len(encabezados) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            if col_idx == 6:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right", vertical="center")

    # Fila de total acumulado
    fila_total = [
        "",
        "",
        "",
        "",
        "TOTAL ACUMULADO (m)",
        round(total_acumulado, 2),
        "",
        "",
        "",
        "",
        "",
    ]
    ws.append(fila_total)
    ultima_fila = ws.max_row

    for col_idx in range(1, len(encabezados) + 1):
        c = ws.cell(row=ultima_fila, column=col_idx)
        c.fill = total_fill
        c.font = total_font
        c.border = thin_border
        if col_idx == 5:
            c.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx == 6:
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right", vertical="center")

    # Si se filtró por una línea específica, agregar fila informativa de meta y % de avance
    if linea_filtrada is not None:
        porcentaje = (
            round((total_acumulado / float(linea_filtrada.metros_totales)) * 100.0, 2)
            if linea_filtrada.metros_totales > 0
            else 0.0
        )
        ws.append(
            [
                "",
                "",
                "",
                "",
                f"META LÍNEA ({linea_filtrada.metros_totales:.2f} m) — AVANCE (%)",
                f"{porcentaje:.2f}%",
                "",
                "",
                "",
                "",
                "",
            ]
        )
        fila_meta = ws.max_row
        for col_idx in range(1, len(encabezados) + 1):
            c = ws.cell(row=fila_meta, column=col_idx)
            c.font = Font(bold=True, size=10)
            c.border = thin_border
            if col_idx in (5, 6):
                c.alignment = Alignment(horizontal="right", vertical="center")

    # Ajustar ancho de columnas automáticamente
    for col_idx in range(1, len(encabezados) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    sufijo = f"_linea_{linea_id}" if linea_id is not None else "_general"
    nombre_archivo = f"reporte_avance{sufijo}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )